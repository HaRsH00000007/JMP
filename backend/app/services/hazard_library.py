"""Hazard-library ingestion (Excel → DB) and read access.

The Excel file is authoritative. Ingestion stores every value verbatim, derives nothing it cannot
justify, and refuses to load a file whose values are internally inconsistent (e.g. an RPN that is not
probability+severity) or whose risk-matrix image differs from the transcribed configuration.
"""

from __future__ import annotations

import hashlib
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from sqlalchemy import select, update
from sqlalchemy.orm import Session, selectinload

from app import rules_config
from app.db.models import Hazard, HazardLibraryVersion, RiskMatrixCell
from app.errors import ErrorCode, JmpError

EXPECTED_HEADERS = {
    "A": "Sr No",
    "B": "Type of Hazard on the road",
    "D": "Severity",
    "E": "Probability",
    "F": "RPN (SXP)",
}
_URL_RE = re.compile(r"https?://\S+")


class IngestionError(JmpError):
    code = ErrorCode.HAZARD_LIBRARY_MISSING


@dataclass(frozen=True)
class HazardDef:
    """Immutable in-memory view of one library hazard (used by the engine, LLM prefix and renderer)."""

    id: str
    code: str
    sr_no: int
    name: str
    severity: int
    probability: str
    rpn_code: str
    severity_band: str
    matrix_zone: str
    control_2w_items: tuple[str, ...]
    control_4w_items: tuple[str, ...]
    control_2w_raw: str
    control_4w_raw: str
    control_short_2w: str | None = None
    control_short_4w: str | None = None

    def items_for(self, vehicle: str) -> tuple[str, ...]:
        return self.control_2w_items if vehicle == "2W" else self.control_4w_items


@dataclass(frozen=True)
class HazardLibrary:
    version_id: str
    version: str
    hazards: tuple[HazardDef, ...]
    matrix: dict[tuple[str, int], str]
    header_fields: tuple[str, ...] = ()
    hospital_network_url: str | None = None
    _by_code: dict[str, HazardDef] = field(default_factory=dict, compare=False, repr=False)

    def __post_init__(self) -> None:
        object.__setattr__(self, "_by_code", {h.code: h for h in self.hazards})

    def get(self, code: str) -> HazardDef:
        return self._by_code[code]

    @property
    def codes(self) -> list[str]:
        return [h.code for h in self.hazards]


# ------------------------------------------------------------------------------------------ parsing
def parse_bullets(raw: str) -> list[str]:
    """Split a control cell into bullet items. Text inside each bullet is left exactly as supplied;
    only leading bullet glyphs/whitespace are removed and exact duplicates dropped."""
    parts = re.split(r"\n|•", raw or "")
    items: list[str] = []
    for p in parts:
        t = p.strip()
        if not t:
            continue
        if t not in items:
            items.append(t)
    return items


def _sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def _matrix_image_sha(path: Path) -> str | None:
    with zipfile.ZipFile(path) as z:
        names = sorted(n for n in z.namelist() if n.startswith("xl/media/") and n.endswith(".png"))
        if not names:
            return None
        return hashlib.sha256(z.read(names[0])).hexdigest()


def parse_workbook(path: Path) -> dict[str, Any]:
    """Parse and validate the workbook. Returns plain data; raises IngestionError on any inconsistency."""
    if not path.exists():
        raise IngestionError(f"Hazard workbook not found: {path}")
    wb = openpyxl.load_workbook(path, data_only=False)
    if "JMP Template" not in wb.sheetnames:
        raise IngestionError("Sheet 'JMP Template' missing")
    ws = wb["JMP Template"]
    for col, expected in EXPECTED_HEADERS.items():
        actual = str(ws[f"{col}2"].value or "").strip()
        if actual != expected:
            raise IngestionError(f"Unexpected header in {col}2: {actual!r} (expected {expected!r})")

    hazards: list[dict[str, Any]] = []
    errors: list[str] = []
    for row in range(3, 28):
        sr = ws[f"A{row}"].value
        name = ws[f"B{row}"].value
        sev = ws[f"D{row}"].value
        prob = ws[f"E{row}"].value
        rpn_formula = ws[f"F{row}"].value
        c2 = ws[f"G{row}"].value or ""
        c4 = ws[f"H{row}"].value or ""
        if not isinstance(sr, int) or not name:
            errors.append(f"row {row}: missing Sr No or hazard name")
            continue
        if not isinstance(sev, int) or not 1 <= sev <= 5:
            errors.append(f"row {row}: severity {sev!r} not an integer 1-5")
        if not isinstance(prob, str) or prob.strip() not in {"A", "B", "C", "D", "E"}:
            errors.append(f"row {row}: probability {prob!r} not A-E")
            continue
        prob = prob.strip()
        expected_formula = f"=CONCATENATE(E{row},D{row})"
        if isinstance(rpn_formula, str) and rpn_formula.startswith("="):
            if rpn_formula.replace(" ", "").upper() != expected_formula:
                errors.append(f"row {row}: RPN formula {rpn_formula!r} is not {expected_formula}")
            rpn = f"{prob}{sev}"
        else:
            rpn = str(rpn_formula or "").strip()
            if rpn != f"{prob}{sev}":
                errors.append(f"row {row}: RPN {rpn!r} != probability+severity {prob}{sev}")
        hazards.append({
            "sr_no": sr, "code": f"HZ-{sr:02d}", "name": str(name).strip(), "severity": sev,
            "probability": prob, "rpn_code": rpn, "control_2w_raw": str(c2), "control_4w_raw": str(c4),
        })
    if len(hazards) != 25:
        errors.append(f"expected 25 hazards, found {len(hazards)}")
    if sorted(h["sr_no"] for h in hazards) != list(range(1, 26)):
        errors.append("Sr No values are not exactly 1..25")
    if len({h["name"] for h in hazards}) != len(hazards):
        errors.append("hazard names are not unique")
    if errors:
        raise IngestionError("Hazard workbook failed validation", details=errors)

    header_fields: list[str] = []
    if "Header" in wb.sheetnames:
        hs = wb["Header"]
        for r in range(2, hs.max_row + 1):
            v = hs[f"B{r}"].value
            if v:
                header_fields.append(str(v).strip())
    hospital_url = None
    for r in range(28, min(ws.max_row, 40) + 1):
        for col in "ABCDEFGH":
            v = ws[f"{col}{r}"].value
            if isinstance(v, str) and (m := _URL_RE.search(v)):
                hospital_url = m.group(0)
    return {
        "hazards": hazards,
        "header_fields": header_fields,
        "hospital_network_url": hospital_url,
        "source_sha256": _sha256_file(path),
        "risk_matrix_image_sha256": _matrix_image_sha(path),
    }


# ---------------------------------------------------------------------------------------- ingestion
def ingest(session: Session, path: Path, version: str, *, activate: bool = True) -> HazardLibraryVersion:
    """Idempotent: re-ingesting the same file under the same version is a no-op; a different file under
    an existing version is refused (versions are immutable)."""
    data = parse_workbook(path)
    matrix_cfg = rules_config.risk_matrix()
    expected_img = matrix_cfg.get("source_image_sha256")
    if data["risk_matrix_image_sha256"] and expected_img and data["risk_matrix_image_sha256"] != expected_img:
        raise IngestionError(
            "Risk-matrix image in the workbook differs from the transcribed config/risk_matrix.yaml — "
            "re-transcribe and review before ingesting.",
            details={"workbook": data["risk_matrix_image_sha256"], "config": expected_img},
        )

    existing = session.scalar(select(HazardLibraryVersion).where(HazardLibraryVersion.version == version))
    if existing is not None:
        if existing.source_sha256 != data["source_sha256"]:
            raise IngestionError(
                f"Library version {version} already exists with a different source file; use a new version."
            )
        if activate and not existing.is_active:
            _activate(session, existing)
        return existing

    band_map = {int(k): v for k, v in matrix_cfg["severity_band"].items()}
    cells = matrix_cfg["cells"]
    detectors = rules_config.hazard_rules().get("detectors", {})
    lib = HazardLibraryVersion(
        version=version,
        source_filename=path.name,
        source_sha256=data["source_sha256"],
        risk_matrix_image_sha256=data["risk_matrix_image_sha256"],
        header_fields=data["header_fields"],
        hospital_network_url=data["hospital_network_url"],
        is_active=False,
        notes="Ingested from Excel; S/P/RPN/controls stored verbatim.",
    )
    session.add(lib)
    session.flush()
    for sev_key, row in cells.items():
        for prob, zone in row.items():
            session.add(RiskMatrixCell(library_version_id=lib.id, probability=prob, severity=int(sev_key), zone=zone))
    for h in data["hazards"]:
        session.add(Hazard(
            library_version_id=lib.id,
            code=h["code"],
            sr_no=h["sr_no"],
            name=h["name"],
            severity=h["severity"],
            probability=h["probability"],
            rpn_code=h["rpn_code"],
            severity_band=band_map[h["severity"]],
            matrix_zone=cells[h["severity"]][h["probability"]],
            control_2w_raw=h["control_2w_raw"],
            control_4w_raw=h["control_4w_raw"],
            control_2w_items=parse_bullets(h["control_2w_raw"]),
            control_4w_items=parse_bullets(h["control_4w_raw"]),
            detection_profile=detectors.get(h["code"], {}),
        ))
    session.flush()
    if activate:
        _activate(session, lib)
    _CACHE.clear()
    return lib


def _activate(session: Session, lib: HazardLibraryVersion) -> None:
    session.execute(update(HazardLibraryVersion).values(is_active=False))
    lib.is_active = True
    session.flush()
    _CACHE.clear()


def apply_approved_short_controls(session: Session, lib: HazardLibraryVersion) -> int:
    """Copy EHS-approved one-line controls (config/hazard_display.yaml) onto the library. No-op while
    the file is unapproved (D-07)."""
    cfg = rules_config.hazard_display()
    if not cfg.get("approved"):
        return 0
    n = 0
    for hz in lib.hazards:
        sc = cfg.get("short_controls", {}).get(hz.code)
        if sc:
            hz.control_short_2w = sc.get("2W")
            hz.control_short_4w = sc.get("4W")
            n += 1
    _CACHE.clear()
    return n


# ------------------------------------------------------------------------------------------- access
_CACHE: dict[str, HazardLibrary] = {}


def load_active_library(session: Session) -> HazardLibrary:
    lib = session.scalar(
        select(HazardLibraryVersion)
        .where(HazardLibraryVersion.is_active.is_(True))
        .options(selectinload(HazardLibraryVersion.hazards), selectinload(HazardLibraryVersion.matrix_cells))
    )
    if lib is None:
        raise IngestionError("No active hazard library. Run: python -m app.cli.ingest_hazards")
    key = str(lib.id)
    if key in _CACHE:
        return _CACHE[key]
    result = HazardLibrary(
        version_id=key,
        version=lib.version,
        hazards=tuple(
            HazardDef(
                id=str(h.id), code=h.code, sr_no=h.sr_no, name=h.name, severity=h.severity,
                probability=h.probability, rpn_code=h.rpn_code, severity_band=h.severity_band,
                matrix_zone=h.matrix_zone, control_2w_items=tuple(h.control_2w_items),
                control_4w_items=tuple(h.control_4w_items), control_2w_raw=h.control_2w_raw,
                control_4w_raw=h.control_4w_raw, control_short_2w=h.control_short_2w,
                control_short_4w=h.control_short_4w,
            )
            for h in sorted(lib.hazards, key=lambda x: x.sr_no)
        ),
        matrix={(c.probability, c.severity): c.zone for c in lib.matrix_cells},
        header_fields=tuple(lib.header_fields or ()),
        hospital_network_url=lib.hospital_network_url,
    )
    _CACHE[key] = result
    return result


# ----------------------------------------------------------------------------------- display helpers
_GENERIC_BULLET = re.compile(r"^mandatory use of", re.I)


def display_text(text: str) -> str:
    """Apply EHS-approved typo corrections only when approved AND enabled (D-06); else verbatim."""
    from app.settings import settings

    cfg = rules_config.hazard_display()
    if settings().hazard_display_text != "approved_corrections" or not cfg.get("approved"):
        return text
    for wrong, right in (cfg.get("text_corrections") or {}).items():
        text = text.replace(wrong, right)
    return text


def short_control(h: HazardDef, vehicle: str) -> str:
    """One-line control: EHS-approved short text if present (D-07), otherwise the first two verbatim
    library bullets (the generic seat-belt/safety-gear bullet is omitted here and shown elsewhere)."""
    approved = h.control_short_2w if vehicle == "2W" else h.control_short_4w
    if approved:
        return approved
    items = [i for i in h.items_for(vehicle) if not _GENERIC_BULLET.match(i)]
    return display_text("; ".join(items[:2]))


def control_items(h: HazardDef, vehicle: str) -> list[str]:
    return [display_text(i) for i in h.items_for(vehicle)]
