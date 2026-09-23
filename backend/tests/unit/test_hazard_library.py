"""Hazard ingestion: the 25 Excel values are reproduced exactly; bad inputs are refused."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from app.services import hazard_library
from app.services.hazard_library import IngestionError, parse_bullets, parse_workbook, short_control
from tests.conftest import REPO

XLSX = REPO / "source" / "JMP Template- 25 Hazards.xlsx"

# Independently transcribed from the workbook (Sr No: name, severity, probability)
EXPECTED = {
    1: ("Sharp Turn / Hairpin Bend", 5, "D"), 2: ("Steep Uphill Slope", 4, "D"), 3: ("Downhill Slope", 5, "D"),
    4: ("Narrow Road", 3, "D"), 5: ("Potholes / Broken Surface", 5, "D"), 6: ("Unmarked Speed Breaker", 3, "D"),
    7: ("Bridge / Culvert", 5, "D"), 8: ("Railway Crossing (Unmanned)", 5, "D"), 9: ("Blind Curve", 4, "D"),
    10: ("Poor Road Markings/ No speed limit signage on road", 4, "D"),
    11: ("Roadside Construction / Diversion", 4, "D"), 12: ("Waterlogging / Flooded Patch", 4, "D"),
    13: ("Landslide / Rockfall Zone", 5, "D"), 14: ("Fog-Prone Stretch", 4, "D"), 15: ("Busy Intersection", 4, "D"),
    16: ("Village Road Crossing", 3, "D"), 17: ("Overhead Obstacles (low bridge, wires)", 3, "D"),
    18: ("Slippery Road (oil, wet)/ Muddy Patch", 4, "D"), 19: ("Animal Crossing / Forest roads", 4, "D"),
    20: ("Pedestrian-Prone Area", 3, "D"), 21: ("Ghats / Bridges", 5, "D"), 22: ("Mobile Network Dead Zone", 2, "D"),
    23: ("Poorly Lit Road (Night Travel)", 4, "D"), 24: ("National Highway", 5, "E"), 25: ("Dangerous Dip", 3, "E"),
}


def test_all_25_hazards_verbatim(library):
    assert library.version == "1.0"
    assert len(library.hazards) == 25
    for h in library.hazards:
        name, sev, prob = EXPECTED[h.sr_no]
        assert (h.name, h.severity, h.probability) == (name, sev, prob), h.code
        assert h.rpn_code == f"{prob}{sev}"
        assert h.code == f"HZ-{h.sr_no:02d}"


def test_controls_preserved_per_vehicle_type(library):
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["JMP Template"]
    for h in library.hazards:
        row = h.sr_no + 2
        assert h.control_2w_raw == ws[f"G{row}"].value
        assert h.control_4w_raw == ws[f"H{row}"].value
    rc = library.get("HZ-08")
    assert any("10 ti 20" in i for i in rc.control_2w_items)  # typo kept verbatim (D-06)
    assert any("seat belt" in i for i in rc.control_4w_items)
    assert not any("seat belt" in i for i in rc.control_2w_items)


def test_severity_band_and_matrix_zone(library):
    assert library.get("HZ-08").severity_band == "HIGH"
    assert library.get("HZ-12").severity_band == "MEDIUM"
    assert library.get("HZ-20").severity_band == "LOW"
    # Danone matrix (Excel image): D4 red, D3 yellow, E3 red, D2 yellow
    assert library.get("HZ-12").matrix_zone == "RED"
    assert library.get("HZ-20").matrix_zone == "YELLOW"
    assert library.get("HZ-25").matrix_zone == "RED"
    assert library.get("HZ-22").matrix_zone == "YELLOW"
    zones = [h.matrix_zone for h in library.hazards]
    assert zones.count("RED") == 19 and zones.count("YELLOW") == 6


def test_header_sheet_and_hospital_link(library):
    assert "Nearest Hospital" in library.header_fields
    assert library.hospital_network_url and "paramounttpa.com" in library.hospital_network_url


def test_reingest_same_version_is_idempotent(db):
    lib = hazard_library.ingest(db, XLSX, "1.0")
    assert lib.version == "1.0" and lib.is_active


def test_different_file_under_same_version_refused(db, tmp_path: Path):
    wb = openpyxl.load_workbook(XLSX)
    wb["JMP Template"]["G3"].value = "changed"
    p = tmp_path / "changed.xlsx"
    wb.save(p)
    with pytest.raises(IngestionError, match="different source file"):
        hazard_library.ingest(db, p, "1.0")


def test_inconsistent_rpn_refused(tmp_path: Path):
    wb = openpyxl.load_workbook(XLSX)
    wb["JMP Template"]["F5"].value = "=CONCATENATE(D5,E5)"
    p = tmp_path / "bad.xlsx"
    wb.save(p)
    with pytest.raises(IngestionError) as ei:
        parse_workbook(p)
    assert any("RPN formula" in d for d in ei.value.details)


def test_bad_severity_refused(tmp_path: Path):
    wb = openpyxl.load_workbook(XLSX)
    wb["JMP Template"]["D7"].value = 9
    p = tmp_path / "bad.xlsx"
    wb.save(p)
    with pytest.raises(IngestionError):
        parse_workbook(p)


def test_parse_bullets_keeps_text_and_drops_exact_duplicates():
    raw = "• Watch for wires \n•Mandatory use of front and rear seat belt\n•Mandatory use of front and rear seat belt"
    assert parse_bullets(raw) == ["Watch for wires", "Mandatory use of front and rear seat belt"]


def test_short_control_uses_verbatim_bullets_until_approved(library):
    s = short_control(library.get("HZ-05"), "4W")
    assert s == "Slow down; Maintain control"
    assert "Mandatory" not in s
